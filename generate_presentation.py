import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # Set slide dimensions to 16:9 widescreen (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    DARK_BLUE = RGBColor(15, 37, 55)       # Primary Dark (#0F2537)
    PRIMARY_TEAL = RGBColor(13, 148, 136)   # Medical Teal (#0D9488)
    LIGHT_BG = RGBColor(248, 250, 252)     # Off-white / Slate (#F8FAFC)
    CARD_BG = RGBColor(255, 255, 255)      # Pure White
    CARD_BORDER = RGBColor(226, 232, 240)  # Light Border (#E2E8F0)
    TEXT_DARK = RGBColor(30, 41, 59)       # Slate 800 (#1E293B)
    TEXT_MUTED = RGBColor(100, 116, 139)   # Slate 500 (#64748B)
    ACCENT_BLUE = RGBColor(37, 99, 235)    # Blue 600 (#2563EB)
    ACCENT_GREEN = RGBColor(22, 163, 74)   # Green 600 (#16A34A)
    ACCENT_AMBER = RGBColor(217, 119, 6)   # Amber 600 (#D97706)

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="MED_FINDER UPDATE"):
        # Header container
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Category pill/subtitle
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PRIMARY_TEAL

        # Main Title
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = DARK_BLUE
        p_title.space_before = Pt(4)

    def set_slide_bg(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide1, DARK_BLUE)

    # Accent decorative bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.8), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY_TEAL
    bar.line.fill.background()

    # Title box
    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.5), Inches(4.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "MED_FINDER"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    p2 = tf1.add_paragraph()
    p2.text = "AI-Powered Medicine Search, Real-Time Pricing & Guidance Platform"
    p2.font.size = Pt(22)
    p2.font.color.rgb = PRIMARY_TEAL
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Project Status, Architecture & Major Updates Summary"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(20)

    p4 = tf1.add_paragraph()
    p4.text = "Key Highlights: RxNorm Global API Fallback | PharmEasy Real-Time Scraper | Intelligent Caching | OCR & Substitute Engine"
    p4.font.size = Pt(12)
    p4.font.color.rgb = RGBColor(148, 163, 184)
    p4.space_before = Pt(24)

    # -------------------------------------------------------------
    # SLIDE 2: Executive Summary & What's New
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide2, LIGHT_BG)
    add_header(slide2, "Executive Summary & Latest Release Updates", "Release Overview")

    cards_data_s2 = [
        ("1. RxNorm API Integration", "Global Catalog Fallback",
         "When a searched medicine is absent from the local DB, the system queries RxNorm (NIH API), formats dosage/strength/category, and saves it seamlessly for instant future access.",
         PRIMARY_TEAL, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.4)),
        ("2. Real-Time Price Scraping", "PharmEasy Live Price Feed",
         "Built a resilient price scraper targeting PharmEasy Next.js hydration payload. Auto-extracts live MRP and discounted pricing for medicines retrieved from external APIs.",
         ACCENT_BLUE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.4)),
        ("3. Duplicate Check & Caching", "Robust Data Integrity",
         "Introduced pre-insert verification in `save_api_medicine()`. Prevents duplicate API entries, minimizes DB bloat, and provides idempotent data caching.",
         ACCENT_GREEN, Inches(0.8), Inches(4.5), Inches(5.6), Inches(2.4)),
        ("4. Cross-Platform Launchers", "Zero-Config Startup",
         "Added `start.bat` and `start.sh` automation scripts for Windows and Unix. Handles virtualenv checks, dependencies, database setup, and launches the Flask server.",
         ACCENT_AMBER, Inches(6.8), Inches(4.5), Inches(5.7), Inches(2.4))
    ]

    for title, subtitle, desc, accent_color, x, y, w, h in cards_data_s2:
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Top accent strip
        strip = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.08))
        strip.fill.solid()
        strip.fill.fore_color.rgb = accent_color
        strip.line.fill.background()

        tb = slide2.shapes.add_textbox(x + Inches(0.2), y + Inches(0.18), w - Inches(0.4), h - Inches(0.3))
        ctf = tb.text_frame
        ctf.word_wrap = True
        ctf.margin_top = ctf.margin_left = ctf.margin_right = ctf.margin_bottom = 0

        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(15)
        cp1.font.bold = True
        cp1.font.color.rgb = DARK_BLUE

        cp2 = ctf.add_paragraph()
        cp2.text = subtitle.upper()
        cp2.font.size = Pt(9.5)
        cp2.font.bold = True
        cp2.font.color.rgb = accent_color
        cp2.space_before = Pt(2)

        cp3 = ctf.add_paragraph()
        cp3.text = desc
        cp3.font.size = Pt(11)
        cp3.font.color.rgb = TEXT_DARK
        cp3.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 3: Deep Dive: RxNorm API Integration
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide3, LIGHT_BG)
    add_header(slide3, "Deep Dive: RxNorm External API Integration", "Core Feature Update")

    # Left Column (Workflow)
    c_left = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1))
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = CARD_BG
    c_left.line.color.rgb = CARD_BORDER

    tb_l = slide3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.7))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "How External Search Fallback Works"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    steps = [
        ("Step 1: Local DB Query", "User searches a medicine. If local SQLite has < 3 results, external search triggers."),
        ("Step 2: RxNorm REST API Call", "Calls NIH RxNav REST endpoint with timeout protection (5s). Retrieves matching active concepts."),
        ("Step 3: Clinical Parsing", "Extracts strength (mg/ml), dosage form (Tablet, Capsule, Syrup), and standardized chemical composition."),
        ("Step 4: Real-Time Price Enrichment", "Asynchronously fetches current Indian market pricing via PharmEasy scraper."),
        ("Step 5: Cache & Store", "Saves newly discovered medicines with source='api' into local DB with duplicate validation.")
    ]

    for st_title, st_desc in steps:
        p_st = tf_l.add_paragraph()
        p_st.text = f"• {st_title}"
        p_st.font.size = Pt(11.5)
        p_st.font.bold = True
        p_st.font.color.rgb = PRIMARY_TEAL
        p_st.space_before = Pt(8)

        p_std = tf_l.add_paragraph()
        p_std.text = f"   {st_desc}"
        p_std.font.size = Pt(10)
        p_std.font.color.rgb = TEXT_DARK

    # Right Column (Benefits & Technical Specs)
    c_right = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1))
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = CARD_BG
    c_right.line.color.rgb = CARD_BORDER

    tb_r = slide3.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.7))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "Technical Implementation & Impact"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    specs = [
        ("Zero 'Not Found' Results", "Expands medicine coverage from thousands of local entries to 100,000+ global medicines."),
        ("Normalized Data Structure", "Converts external RxNorm payloads into MED_FINDER's unified relational schema."),
        ("Configurable Parameters", "Configured in config.py: RXNORM_API_BASE_URL, RXNORM_API_TIMEOUT (5s)."),
        ("Source Attribution", "Flags medicines with source='api' vs 'dataset' for clear visual badge & data traceability in UI."),
        ("Self-Growing Database", "The system enriches its own database progressively as users search new medicines.")
    ]

    for sp_title, sp_desc in specs:
        p_sp = tf_r.add_paragraph()
        p_sp.text = f"✔ {sp_title}"
        p_sp.font.size = Pt(11.5)
        p_sp.font.bold = True
        p_sp.font.color.rgb = ACCENT_BLUE
        p_sp.space_before = Pt(8)

        p_spd = tf_r.add_paragraph()
        p_spd.text = f"   {sp_desc}"
        p_spd.font.size = Pt(10)
        p_spd.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 4: Deep Dive: PharmEasy Real-Time Price Scraper
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide4, LIGHT_BG)
    add_header(slide4, "Deep Dive: Real-Time Price Scraping Engine", "Pricing Service Update")

    p_cards = [
        ("HTML / Next.js Hydration Extraction", "Advanced Scraper Architecture",
         "Instead of fragile DOM element scraping, the engine parses `<script id='__NEXT_DATA__'>` JSON state from PharmEasy search results. Extracts precise numeric MRP & discounted prices without brittle CSS selectors.",
         PRIMARY_TEAL, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.1)),
        ("Multi-Tier Fallback & Sanitization", "Resilient Scraping Logic",
         "1. Primary: JSON Next.js state extraction\n2. Secondary: Regex pattern extraction on product card HTML\n3. Tertiary: Fallback to estimated generic pricing.\nFeatures 3-second timeout protection and realistic User-Agent headers.",
         ACCENT_BLUE, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.1)),
        ("Price Backfill & Batch Updating", "Database Price Enrichment",
         "Includes standalone maintenance utilities to scan medicines with missing/zero prices and auto-backfill verified Indian market pricing (₹ INR), ensuring users always see realistic cost comparisons.",
         ACCENT_GREEN, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.1))
    ]

    for title, subtitle, desc, col, x, y, w, h in p_cards:
        c = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER

        st = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.08))
        st.fill.solid()
        st.fill.fore_color.rgb = col
        st.line.fill.background()

        tb = slide4.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), w - Inches(0.4), h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = DARK_BLUE

        p2 = tf.add_paragraph()
        p2.text = subtitle.upper()
        p2.font.size = Pt(9)
        p2.font.bold = True
        p2.font.color.rgb = col
        p2.space_before = Pt(3)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = TEXT_DARK
        p3.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 5: Full Platform Feature Matrix
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide5, LIGHT_BG)
    add_header(slide5, "MED_FINDER Complete Platform Capabilities", "System Features")

    feat_matrix = [
        ("Smart Medicine Search", "Instant autocomplete & fuzzy search across brand names, active salts & manufacturers.", Inches(0.8), Inches(1.8)),
        ("Generic Substitutes Engine", "Recommends chemically identical generic alternatives to help patients save 50-80% on costs.", Inches(4.8), Inches(1.8)),
        ("Prescription OCR Scanner", "Upload prescription images or medicine strip photos to extract and identify medicine names.", Inches(8.8), Inches(1.8)),
        ("Disease & Symptom Guidance", "Symptom checker providing disease overviews, specialist recommendations & first-line remedies.", Inches(0.8), Inches(4.5)),
        ("AI Health Assistant", "Interactive assistant answering user questions regarding dosage precautions & drug uses.", Inches(4.8), Inches(4.5)),
        ("Cross-Platform Deployable", "Self-contained SQLite architecture with zero external database dependencies + automated scripts.", Inches(8.8), Inches(4.5))
    ]

    for f_title, f_desc, x, y in feat_matrix:
        c = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.6), Inches(2.4))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER

        tb = slide5.shapes.add_textbox(x + Inches(0.2), y + Inches(0.18), Inches(3.2), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = f_title
        p1.font.size = Pt(13.5)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_TEAL

        p2 = tf.add_paragraph()
        p2.text = f_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 6: Architecture & Data Flow
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide6, LIGHT_BG)
    add_header(slide6, "System Architecture & Request Pipeline", "Technical Architecture")

    arch_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1))
    arch_box.fill.solid()
    arch_box.fill.fore_color.rgb = CARD_BG
    arch_box.line.color.rgb = CARD_BORDER

    tb_arch = slide6.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.6))
    tf_a = tb_arch.text_frame
    tf_a.word_wrap = True

    p = tf_a.paragraphs[0]
    p.text = "Modular Component Flow"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    layers = [
        ("1. Presentation Layer (Frontend)", "Jinja2 HTML Templates, Bootstrap 5 UI, Vanilla JS Autocomplete, Responsive Mobile-First Design"),
        ("2. Application Layer (Flask Routes)", "REST Endpoints: `/search`, `/api/autocomplete`, `/medicine/<id>`, `/scan`, `/disease-guidance`, `/assistant`"),
        ("3. Service Layer (Business Logic)", "• medicine_service.py (Search & Substitutes)   • external_api_service.py (RxNorm Fallback)\n• pharmeasy_service.py (Price Scraper)       • ocr_service.py (Prescription Recognition)\n• disease_service.py (Symptom Matching)     • recommendation_service.py (Substitutes)"),
        ("4. Persistence Layer (SQLite Database)", "Schema: `medicines` (name, salt, strength, price, source), `disease_guidance`, `manufacturers`.\nIndexed for fast substring & full-text searches with zero external server dependencies.")
    ]

    for l_title, l_desc in layers:
        p_t = tf_a.add_paragraph()
        p_t.text = l_title
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = PRIMARY_TEAL
        p_t.space_before = Pt(8)

        p_d = tf_a.add_paragraph()
        p_d.text = l_desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 7: Roadmap & Next Steps
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide7, LIGHT_BG)
    add_header(slide7, "Impact, Results & Future Roadmap", "Project Value & Future")

    r_cards = [
        ("Current Milestones Achieved",
         "• 100% search hit rate with RxNorm API fallback\n• Automated live price fetching with PharmEasy scraper\n• Clean idempotent DB persistence with duplicate checks\n• End-to-end OCR prescription recognition\n• Cross-platform single-click launch scripts",
         ACCENT_GREEN, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1)),
        ("Next Phase Roadmap",
         "• Multi-Pharmacy Price Comparison (1mg, Netmeds, Apollo)\n• Pharmacy Geo-Location & Nearest Availability mapping\n• LLM-Powered Prescription Simplifier (Plain English translation)\n• Offline PWA Mobile App Support\n• Drug-Drug Interaction Warning Matrix",
         ACCENT_BLUE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1))
    ]

    for title, desc, col, x, y, w, h in r_cards:
        c = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER

        st = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.08))
        st.fill.solid()
        st.fill.fore_color.rgb = col
        st.line.fill.background()

        tb = slide7.shapes.add_textbox(x + Inches(0.3), y + Inches(0.3), w - Inches(0.6), h - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = DARK_BLUE

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(14)

    output_path = Path("reports/MED_FINDER_Project_Update.pptx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"Presentation saved successfully to: {output_path.resolve()}")

if __name__ == "__main__":
    create_deck()
